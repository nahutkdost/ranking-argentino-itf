#!/usr/bin/env python3
"""
Generador del sitio "Ranking Clasificatorio — Mundial ITF Paraguay 2027".

Uso:
    python3 build.py Ranking_General.xlsx            -> genera index.html
    python3 build.py Ranking_General.xlsx -o otro.html

Después de cada fecha:
  1. Cargá los puntajes en la planilla (columnas "1er/2do/3er Ranqueable") y guardala.
  2. Si cambian datos de un evento (fecha, sede, flyer), editá EVENTOS más abajo.
  3. Corré este script y subí el index.html resultante.

Requiere: pip install openpyxl
"""
import base64, json, re, sys, datetime, unicodedata, argparse
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

HERE = Path(__file__).resolve().parent

# ---------------------------------------------------------------------------
# EVENTOS DEL CALENDARIO CLASIFICATORIO  (editar acá cuando se confirmen datos)
# 'flyer': ruta a una imagen (jpg/png) dentro de assets/, o None.
# El estado (disputado / próximo / a confirmar) se calcula solo: un evento
# figura "disputado" cuando la planilla tiene puntajes cargados en su columna.
# ---------------------------------------------------------------------------
EVENTOS = [
    {
        "id": "arg", "corto": "ARG", "nombre": "Campeonato Argentino 2026",
        "orden": "Fecha 1 de 4",
        "fecha": "2026-07-25", "fecha_fin": None, "fecha_texto": "25 de julio de 2026",
        "sede": "Estadio MEPA", "ciudad": "Buenos Aires",
        "detalle": "Campeonato Nacional de Taekwon-Do ITF organizado por FETRA. Primera fecha puntuable del ranking clasificatorio al Mundial Paraguay 2027.",
        "flyer": "assets/flyer-argentino.jpg",
        "direccion": None,
        "puntos": [6, 4, 3, 2],
        "cronograma": [], "secciones": [],
    },
    {
        "id": "r1", "corto": "R1", "nombre": "Primer Ranqueable Argentino",
        "orden": "Fecha 2 de 4",
        "fecha": "2026-11-14", "fecha_fin": "2026-11-15", "fecha_texto": "14 y 15 de noviembre de 2026",
        "sede": "Polideportivo “Finito” Gehrmann", "ciudad": "Posadas, Misiones",
        "detalle": "Primer evento ranqueable del proceso clasificatorio al Mundial Paraguay 2027. Fiscaliza FETRA.",
        "flyer": "assets/flyer-ranqueable1.jpg",
        "direccion": "Av. Cmte. Rosales 10, Posadas, Misiones",
        "reglamento": "assets/Reglamento - 1er Ranqueable Paraguay 2027.pdf",
        "puntos": [9, 6, 4, 3, 2, 1],
        "cronograma": [
            {"dia": "Viernes 13 de noviembre", "items": [
                "Acreditación y pesaje en el estadio: 9 a 13 h y 15 a 20 h",
                "Reunión de umpires", "Reunión de coaches"]},
            {"dia": "Sábado 14 de noviembre", "items": ["Competencia Adultos (todas las modalidades)"]},
            {"dia": "Domingo 15 de noviembre", "items": ["Competencia Pre-Junior y Junior (todas las modalidades)"]},
        ],
        "secciones": [
            {"t": "Modalidades", "items": ["Lucha", "Formas", "Rotura de Poder", "Técnicas Especiales", "Lucha Preestablecida"]},
            {"t": "Divisiones por edad", "items": [
                "Pre-Junior: nacidos en 2012, 2013 y 2014",
                "Junior: nacidos en 2009, 2010 y 2011",
                "Adultos: nacidos en 2008 o antes"]},
            {"t": "Graduaciones permitidas", "items": [
                "Cinturones negros de I a VI Dan.",
                "Solo en este ranqueable pueden participar cinturones 1er Gup, excepto en Formas.",
                "Para el segundo ranqueable es obligatorio presentarse con el I Dan aprobado."]},
            {"t": "Pesaje oficial", "items": [
                "Deben presentarse todos los inscriptos en Lucha, con la acreditación realizada.",
                "Si no dan el peso en el primer intento (tolerancia de 500 g), tienen una segunda oportunidad dentro de un máximo de una hora.",
                "Si no cumplen con el peso quedan descalificados, sin posibilidad de pasar a otra categoría de peso."]},
            {"t": "Conservación de puntos", "items": [
                "Quienes sumaron puntos en el Campeonato Argentino 2026 deben inscribirse en la misma categoría para conservar todo su puntaje.",
                "Excepcionalmente, y solo en este ranqueable, quien cambie de categoría puede trasladar a la nueva el 50 % de los puntos obtenidos en la anterior."]},
            {"t": "Asignación de puestos", "items": [
                "Se disputa una ronda adicional para definir el 3° y el 4° puesto.",
                "5° puesto: el eliminado en cuartos de final por el campeón de la categoría.",
                "6° puesto: el eliminado en cuartos de final por el subcampeón de la categoría."]},
        ],
    },
    {
        "id": "r2", "corto": "R2", "nombre": "Segundo Ranqueable Argentino",
        "orden": "Fecha 3 de 4",
        "fecha": None, "fecha_fin": None, "fecha_texto": "Marzo 2027 (fecha estimada)",
        "fecha_corta": "Mar 2027", "estimada": True,
        "sede": "Sede a confirmar", "ciudad": "Buenos Aires",
        "detalle": "Fecha aproximada. Sede y fecha exacta a confirmar por FETRA.",
        "flyer": None,
        "direccion": None, "puntos": [9, 6, 4, 3, 2, 1], "cronograma": [], "secciones": [],
    },
    {
        "id": "r3", "corto": "R3", "nombre": "Tercer Ranqueable Argentino",
        "orden": "Fecha 4 de 4",
        "fecha": None, "fecha_fin": None, "fecha_texto": "Abril / mayo 2027 (fecha estimada)",
        "fecha_corta": "Abr/May 2027", "estimada": True,
        "sede": "Sede a confirmar", "ciudad": "Buenos Aires",
        "detalle": "Fecha aproximada. Última fecha del ranking clasificatorio. Sede y fecha exacta a confirmar por FETRA.",
        "flyer": None,
        "direccion": None, "puntos": [9, 6, 4, 3, 2, 1], "cronograma": [], "secciones": [],
    },
]

# Puntos máximos en juego por cada fecha restante (1° puesto de un ranqueable = 9 pts,
# según el reglamento del 1er Ranqueable). Se usa para la clasificación anticipada.
PTS_MAX_POR_FECHA = 9

# ---------------------------------------------------------------------------
# CLASIFICADOS A MUNDIALES ANTERIORES
# CSV con columnas: codigo;nombre;categoria_original  (nombre como "Apellido Nombre")
# Equipos: "ARGENTINA(Apellido1, Apellido2)". Se vinculan solos a los competidores
# actuales por nombre; el build muestra un reporte para revisar las coincidencias.
# ---------------------------------------------------------------------------
HISTORIAL = [
    {"id": "wc2025", "nombre": "Mundial 2025", "sede": "Poreč, Croacia", "anio": 2025,
     "csv": "assets/clasificados-mundial-2025.csv",
     # Ranking clasificatorio completo de ese Mundial (opcional). Bloques "CÓDIGO título" + filas
     # Nombre | Escuela/Instructor | fechas... | Total | Notas
     "ranking": "assets/ranking-clasificatorio-mundial-2025.xlsx",
     "ranking_eventos": ["CA", "1° Sel.", "2° Sel.", "3° Sel.", "4° Sel."]},
    {"id": "wc2023", "nombre": "Mundial 2023", "sede": "Tampere, Finlandia", "anio": 2023,
     "csv": "assets/participantes-mundial-2023.csv"},
    # 2019: las categorías eran distintas (sin Pre-Juveniles); NO se vinculan a las categorías actuales
    {"id": "wc2019", "nombre": "Mundial 2019", "sede": "Inzell, Alemania", "anio": 2019,
     "csv": "assets/participantes-mundial-2019.csv", "equiv": False},
]
TEAM_CODES = {"53570": "Pre-Juveniles", "54570": "Juveniles", "55580": "Adultos"}


def cat_es(orig):
    """Traduce el nombre ITF en inglés de una categoría a la forma usada en el sitio."""
    o = orig.replace("‐", "-")
    gen = "Masculino" if re.search(r"\bmale\b", o, re.I) else "Femenino" if re.search(r"\bfemale\b", o, re.I) else ""
    age = "Pre-Juveniles" if re.search(r"pre-?junior", o, re.I) else "Juveniles" if re.search(r"\bjunior\b", o, re.I) else "Adultos"
    if re.search(r"pre-?arranged", o, re.I):
        return "Lucha Pre-Establecida", f"Lucha Pre-Establecida {age}"
    if re.search(r"pattern", o, re.I):
        mod = "Formas"
        m = re.search(r"pattern\s+(.+?)\s*dan", o, re.I)
        nums = re.findall(r"\d", m.group(1)) if m else []
        rom = {"1": "I", "2": "II", "3": "III", "4": "IV", "5": "V", "6": "VI"}
        div = ("-".join([rom[nums[0]], rom[nums[-1]]]) if len(nums) > 1 else rom.get(nums[0], "") if nums else "") + " Dan"
    elif re.search(r"special", o, re.I):
        mod, div = "Técnicas Especiales", ""
    elif re.search(r"power", o, re.I):
        mod, div = "Rotura de Poder", ""
    else:
        mod = "Lucha"
        m1 = re.search(r"up to (\d+)", o, re.I); m2 = re.search(r"over \d+ to (\d+)", o, re.I); m3 = re.search(r"over (\d+)\s*kg", o, re.I)
        div = f"-{m1.group(1)} kg" if m1 else f"-{m2.group(1)} kg" if m2 else f"+{m3.group(1)} kg" if m3 else ""
    return mod, " ".join(x for x in (mod, age, gen, div) if x)

SHEETS = [
    ("Formas Adultos", "Formas"), ("Formas Juveniles", "Formas"), ("Formas Pre-Juveniles", "Formas"),
    ("Lucha Adultos", "Lucha"), ("Lucha Juveniles", "Lucha"), ("Lucha Pre-Juveniles", "Lucha"),
    ("Rotura de Poder", "Rotura de Poder"), ("Tecnicas Especiales", "Técnicas Especiales"),
    ("Lucha Pre-Establecida", "Lucha Pre-Establecida"),
]
AGES = ["Pre-Juveniles", "Juveniles", "Adultos"]


def clean_name(s):
    s = re.sub(r"\s+", " ", str(s)).strip()
    s = re.sub(r"\s*,\s*", ", ", s)
    return s


def num(v):
    if v is None or v == "" or v == "-":
        return None
    try:
        return float(v) if not float(v).is_integer() else int(v)
    except (TypeError, ValueError):
        return None


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    cats = []
    for sheet, mod in SHEETS:
        ws = wb[sheet]
        team = mod == "Lucha Pre-Establecida"
        off = 1 if team else 0
        cur = None
        for row in ws.iter_rows(values_only=True):
            row = list(row) + [None] * 20
            a = row[0]
            # Encabezado de categoría
            if isinstance(a, str) and a not in ("Pos", "-") and row[4] is None and row[2] is None:
                m = re.match(r"^\s*(\d+)\s*-\s*(.+)$", a)
                numero, titulo = (m.group(1), m.group(2).strip()) if m else ("", a.strip())
                titulo = titulo.replace("Forma ", "Formas ", 1) if titulo.startswith("Forma ") else titulo
                titulo = titulo.replace("Tecnicas", "Técnicas")
                age = next((x for x in AGES if f" {x}" in f" {titulo}"), "")
                # 'Juveniles' también matchea en 'Pre-Juveniles'
                if "Pre-Juveniles" in titulo:
                    age = "Pre-Juveniles"
                gender = "Femenino" if "Femenino" in titulo else ("Masculino" if "Masculino" in titulo else "Mixto")
                div = titulo.split(gender, 1)[1].strip() if gender in titulo else ""
                cur = {"id": None, "num": numero, "mod": mod, "age": age, "gen": gender, "div": div,
                       "title": titulo, "cupos": 1 if team else 2, "team": team, "rows": []}
                cats.append(cur)
                continue
            if cur is None or row[2] in (None, "") or not isinstance(row[1], (int, float)):
                continue
            name = row[4]
            if not name:
                continue
            cur["id"] = cur["id"] or str(row[2]).strip()
            pts = [num(row[6 + off + i]) for i in range(4)]
            r = {"n": clean_name(name) if not team else str(name).strip(),
                 "s": str(row[5 + off] or "").strip(),
                 "p": pts}
            tie = num(row[10 + off])
            if tie:
                r["d"] = tie
            if team:
                r["m"] = [clean_name(x) for x in str(row[5] or "").split("/") if x.strip()]
            cur["rows"].append(r)
    cats = [c for c in cats if c["rows"]]
    for c in cats:  # el código interno de la planilla (ej. LMAM87) no se publica
        c.pop("id", None)
    return cats


def nrm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ,]+", " ", s)


def slug(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def close(a, b):
    if a == b:
        return True
    if len(a) >= 4 and len(b) >= 4 and (a.startswith(b) or b.startswith(a)):
        return True
    if abs(len(a) - len(b)) > 1 or min(len(a), len(b)) < 4:
        return False
    # distancia de edición <= 1
    i = 0
    while i < min(len(a), len(b)) and a[i] == b[i]:
        i += 1
    return a[i + 1:] == b[i + 1:] or a[i + 1:] == b[i:] or a[i:] == b[i + 1:]


def historial(cats):
    import csv
    people = {}  # nombre normalizado -> {"n", "keys":set, "S":[], "G":[]}
    for c in cats:
        for r in c["rows"]:
            for name in (r.get("m") or [r["n"]]):
                k = nrm(name).strip()
                o = people.setdefault(k, {"n": name, "keys": set()})
                o["keys"].add(slug(name) + "--" + slug(r["s"]))
                sur, _, giv = name.partition(",")
                o["S"], o["G"] = nrm(sur).split(), nrm(giv).split()
    out = []
    for h in HISTORIAL:
        path = HERE / h["csv"]
        if not path.exists():
            print(f"  aviso: no se encontró {h['csv']}")
            continue
        rows = list(csv.DictReader(path.open(encoding="utf-8"), delimiter=";"))
        items, nmatch = [], 0
        indiv_matched = {}  # nombre -> set de grupos de edad (2° dígito del código)
        pending_teams = []
        for r in rows:
            cod, nombre = r["codigo"].strip(), r["nombre"].strip()
            it = {"cod": cod, "orig": r.get("categoria_original", "").strip(), "keys": []}
            it["mod"], it["cat"] = cat_es(it["orig"])
            mt = re.match(r"^\s*ARGENTINA\s*\((.+)\)\s*$", nombre, re.I)
            if mt:
                it["team"] = True
                it["miembros"] = [x.strip() for x in mt.group(1).split(",")]
                it["n"] = " / ".join(it["miembros"])
                pending_teams.append(it)
            else:
                it["n"] = " ".join(w if not w.isupper() else w.title() for w in nombre.split())
                T = nrm(nombre).split()
                cands = []
                for k, o in people.items():
                    if not o["S"] or not o["G"]:
                        continue
                    strong = all(x in T for x in o["S"]) and o["G"][0] in T
                    medium = o["S"][0] == T[0] and any(close(o["G"][0], t) for t in T[1:])
                    if strong or medium:
                        cands.append((0 if strong else 1, k))
                if cands:
                    best = min(x[0] for x in cands)
                    for _, k in [x for x in cands if x[0] == best]:
                        it["keys"] += sorted(people[k]["keys"])
                        indiv_matched.setdefault(k, set()).add(cod[1:2])
                    nmatch += 1
            items.append(it)
        # equipos: por apellido, priorizando a quienes ya figuran como clasificados individuales
        for it in pending_teams:
            for sur in it["miembros"]:
                S = nrm(sur).split()
                cands = [k for k, o in people.items() if o.get("S") == S]
                age = it["cod"][1:2]
                pref = [k for k in cands if age in indiv_matched.get(k, ())] or [k for k in cands if k in indiv_matched]
                use = pref if pref else (cands if len(cands) == 1 else [])
                for k in use:
                    it["keys"] += sorted(people[k]["keys"])
            if it["keys"]:
                nmatch += 1
        print(f"  {h['nombre']}: {len(items)} plazas, {nmatch} vinculadas a competidores actuales")
        rk = prev_ranking(h, people, items) if h.get("ranking") else None
        rep_lines = [f"{i['cod']};{i['n']};" + ",".join(i["keys"]) for i in items]
        (HERE / f"reporte-{h['id']}.txt").write_text("\n".join(rep_lines), encoding="utf-8")
        out.append({k: h[k] for k in ("id", "nombre", "sede", "anio")} | {"items": items, "teamCodes": TEAM_CODES,
                    "equiv": h.get("equiv", True), "clasificados": bool(h.get("ranking")),
                    "ranking": rk, "rankingEventos": h.get("ranking_eventos", [])})
    return out


def match_person(nombre, people):
    """Devuelve claves de competidores actuales para un nombre escrito 'Apellido Nombre'."""
    T = nrm(nombre).split()
    if not T:
        return []
    cands = []
    for k, o in people.items():
        if not o.get("S") or not o.get("G"):
            continue
        strong = all(x in T for x in o["S"]) and o["G"][0] in T
        medium = o["S"][0] == T[0] and any(close(o["G"][0], t) for t in T[1:])
        if strong or medium:
            cands.append((0 if strong else 1, k))
    if not cands:
        return []
    best = min(x[0] for x in cands)
    return sorted({key for _, k in cands if _ == best for key in people[k]["keys"]})


def same_person(a, b):
    A, B = nrm(a).split(), nrm(b).split()
    if not A or not B or A[0] != B[0]:
        return False
    return any(close(x, y) for x in A[1:] for y in B[1:]) or (len(A) == 1 or len(B) == 1)


def map_code(cod):
    # En 2025 la lucha adultos usaba 15171-15177 / 25171-25177; hoy 15181-15187 / 25181-25187
    if re.match(r"^[12]517\d$", cod):
        return cod[:3] + "8" + cod[4]
    return cod


def tcase(x):
    small = {"de", "del", "la", "las", "los", "y", "e"}
    return " ".join(w if i and w.lower() in small and not w.lower().__eq__(w) and False else (w.lower() if i and w.lower() in small else w.capitalize()) for i, w in enumerate(x.lower().split()))


def prev_ranking(h, people, items):
    path = HERE / h["ranking"]
    if not path.exists():
        print(f"  aviso: no se encontró {h['ranking']}")
        return None
    wb = openpyxl.load_workbook(path, data_only=True)
    nev = len(h.get("ranking_eventos", [])) or 5
    blocks, cur = [], None
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            row = list(row) + [None] * 12
            a = row[0]
            m = re.match(r"^\s*(\d{5})\s+(.+)$", str(a or ""))
            if m:
                cur = {"cod": map_code(m.group(1)), "orig": m.group(2).strip(), "rows": []}
                blocks.append(cur)
                continue
            if cur is None or not a or not str(a).strip():
                continue
            pts = [num(row[2 + i]) for i in range(nev)]
            tot = num(row[2 + nev])
            if tot is None:
                tot = sum(x or 0 for x in pts)
            esc = re.sub(r"\s*ARG\s*-\s*ARGENTINA\s*$", "", str(row[1] or "").strip(), flags=re.I).strip()
            nota = str(row[3 + nev] or "").strip()
            raw = re.sub(r"\s+", " ", str(a)).strip()
            r = {"p": pts, "t": tot, "e": tcase(esc) if esc.isupper() else esc}
            mt = re.match(r"^(.*?)\s*\((.+?)\)?\s*$", raw)
            if cur["cod"] in TEAM_CODES:
                miembros = [x.replace("_", " ").strip() for x in (mt.group(2).split("/") if mt else [])]
                r["n"] = tcase(mt.group(1) if mt else raw)
                r["m"] = [tcase(x) for x in miembros]
                r["keys"] = sorted({k for x in miembros for k in match_person(x, people)})
            else:
                r["n"] = tcase(raw)
                r["keys"] = match_person(raw, people)
            cur["rows"].append(r)
    # posiciones y quiénes viajaron (según la lista de clasificados)
    nlink = ntrav = 0
    for b in blocks:
        b["rows"].sort(key=lambda r: -(r["t"] or 0))
        tots = [r["t"] or 0 for r in b["rows"]]
        trav_items = [it for it in items if map_code(it["cod"]) == b["cod"]]
        for r in b["rows"]:
            r["pos"] = 1 + sum(1 for t in tots if t > (r["t"] or 0)) if (r["t"] or 0) > 0 else None
            if b["cod"] in TEAM_CODES:
                mem = " ".join(r.get("m", []))
                r["wc"] = any(any(nrm(x).split()[0] in nrm(mem).split() for x in it.get("miembros", [])) for it in trav_items)
            else:
                r["wc"] = any(same_person(it["n"], r["n"]) for it in trav_items)
            nlink += bool(r["keys"]); ntrav += r["wc"]
        found = sum(1 for r in b["rows"] if r["wc"])
        if found < len(trav_items):
            b["extra"] = [it["n"] for it in trav_items
                          if not any(r["wc"] and (same_person(it["n"], r["n"]) or it.get("team")) for r in b["rows"])]
    print(f"  Ranking {h['nombre']}: {len(blocks)} categorías, {sum(len(b['rows']) for b in blocks)} filas, "
          f"{nlink} vinculadas al ranking actual, {ntrav} marcadas como clasificadas")
    return blocks


def b64(rel):
    if not rel:
        return None
    p = HERE / rel
    if not p.exists():
        print(f"  aviso: no se encontró {rel}")
        return None
    mime = {".png": "image/png", ".pdf": "application/pdf"}.get(p.suffix.lower(), "image/jpeg")
    return f"data:{mime};base64," + base64.b64encode(p.read_bytes()).decode()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default=str(HERE / "index.html"))
    ap.add_argument("--fecha", help="Fecha de actualización (AAAA-MM-DD). Por defecto, hoy.")
    a = ap.parse_args()

    cats = parse(a.xlsx)
    played = [any(r["p"][i] is not None for c in cats for r in c["rows"]) for i in range(4)]
    eventos = []
    for i, e in enumerate(EVENTOS):
        e = dict(e)
        e["jugado"] = played[i]
        e["flyer"] = b64(e["flyer"])
        # Reglamento en PDF (opcional): se incrusta para que el sitio siga siendo un solo archivo
        reg = e.get("reglamento")
        e["reglamento"] = {"data": b64(reg), "nombre": Path(reg).name} if reg and b64(reg) else None
        eventos.append(e)

    data = {
        "actualizado": a.fecha or datetime.date.today().isoformat(),
        "ptsMaxFecha": PTS_MAX_POR_FECHA,
        "eventos": eventos,
        "categorias": cats,
        "historial": historial(cats),
        "mundial": {"flyer": b64("assets/flyer-mundial.jpg")},
    }
    tpl = (HERE / "template.html").read_text(encoding="utf-8")
    html = tpl.replace("/*__DATA__*/null", json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    html = html.replace("__SEAL__", b64("assets/seal.png") or "")
    Path(a.out).write_text(html, encoding="utf-8")
    n = sum(len(c["rows"]) for c in cats)
    print(f"OK -> {a.out}  ({len(cats)} categorías, {n} inscripciones, fechas disputadas: {sum(played)})")


if __name__ == "__main__":
    main()
